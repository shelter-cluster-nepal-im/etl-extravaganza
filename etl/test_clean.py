#TODO: create sheets in mem, don't import

import unittest
import etl
import clean
from openpyxl import load_workbook

sample_wb = load_workbook('/Users/ewanog/code/git_repos/nepal-earthquake/shelter/etl-extravaganza/clean_test.xlsx', data_only = True)
train = sample_wb.get_sheet_by_name('Training')
ref = sample_wb.get_sheet_by_name('Reference')
date_wb = load_workbook('/Users/ewanog/code/git_repos/nepal-earthquake/shelter/etl-extravaganza/date_test.xlsx', data_only = True)
dist = date_wb.get_sheet_by_name('dates')

class TestClean(unittest.TestCase):
    def test_cur(self):
        clean.algo20(dist, ref)

if __name__ == '__main__':
    unittest.main()


    
