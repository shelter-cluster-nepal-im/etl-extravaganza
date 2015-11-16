import unittest
import etl
import master_db_creation
from sqlalchemy import Column, Integer, String, Float
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy.orm import relationship
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy import create_engine
from openpyxl import load_workbook
from openpyxl import Workbook
from openpyxl.cell import column_index_from_string
import os

Base = declarative_base()
engine = create_engine(os.environ['dbk'])
Base.metadata.bind = engine
DBSession = sessionmaker(bind=engine)
session = DBSession()
Base.metadata.create_all(engine)


class TestDBb(unittest.TestCase):

    def test_check_zero_entries(self):
        locs = {'targeting': 1, 'quantity':2, 'total_hh': 3, 'float': 4}
        s = Workbook().active
        s.append(('val','None','1','None'))
        ret = master_db_creation.check_zero_entries(s.rows[0],locs, Base.metadata.tables['test'])

        self.assertTrue(ret[0].value == 'val')
        self.assertTrue(ret[1].value == '0')
        self.assertTrue(ret[2].value == '1')
        self.assertTrue(ret[3].value == '0')

class test(Base):
    __tablename__ = 'test'
    targeting = Column(String(250), primary_key=True)
    quantity = Column(Integer)
    total_hh = Column(Integer)
    float = Column(Float)




if __name__ == '__main__':
    Base.metadata.create_all(engine)
    unittest.main()
